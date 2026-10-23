"""モッチが PC のファイルを探して読む(要件 F22 1.10、基本設計 Q41・S6-9、規程第12条)。

サイトと同じく、設定で本人が登録したフォルダの中だけを見る。読むだけで、書きかえ・消す・動かすはしない。
- 読まないもの:.env、名前に password・secret・token・credential が入るもの、鍵(.pem・.key など)、
  隠しフォルダ・.git・node_modules の中、段取りさん自身の data/
- 開く(いつものアプリで開く)は確認のあと。.exe・.bat など動くファイルは開かない
- 読んだ中身は、その会話の中で Claude に渡すだけ。DB には残さない
"""
import json
import os
import re
import stat
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

KEY = "agent_folders"
HOME = Path.home()                       # はじめの登録(Workspace・Downloads)の元。テストでは差しかえる
DATA_DIR = Path(__file__).resolve().parent / "data"
SKIP_DIRS = {".git", "node_modules", "__pycache__", "venv", ".venv", "$RECYCLE.BIN", "System Volume Information"}
SECRET_NAME = re.compile(r"password|passwd|secret|token|credential|パスワード|暗証", re.I)
KEY_EXT = {".pem", ".key", ".pfx", ".p12", ".kdbx", ".ppk", ".asc", ".gpg"}
RUN_EXT = {".exe", ".bat", ".cmd", ".com", ".ps1", ".psm1", ".vbs", ".vbe", ".js", ".jse", ".wsf", ".wsh", ".hta", ".msi", ".msp",
           ".scr", ".lnk", ".url", ".reg", ".jar", ".py", ".pyw", ".cpl", ".appref-ms"}
TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".py", ".js", ".ts", ".html", ".css", ".xml", ".yml", ".yaml", ".toml", ".ini",
            ".log", ".sql", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".tex", ".bib"}
DOC_EXT = {".docx", ".xlsx", ".xlsm", ".pptx", ".pdf"}
MAX_CHARS = 12000           # 1回に渡す長さ。長いものは頭から
MAX_BYTES = 30 * 1024 * 1024
SEARCH_LIMIT = 20
SEARCH_SECONDS = 8          # 探すのにかける時間の上限(大きなフォルダで止まらないため)


class LocalFileError(Exception):
    pass


# ---------- 登録したフォルダ(設定の「フォルダ」) ----------
def folders(con):
    r = con.execute("SELECT value FROM settings WHERE key = ?", (KEY,)).fetchone()
    if r:
        return json.loads(r["value"])
    return [str(p) for p in (HOME / "Workspace", HOME / "Downloads") if p.is_dir()]  # はじめは Workspace と Downloads(本人の決定 A)


def _save(con, items):
    con.execute("INSERT OR REPLACE INTO settings(key, value) VALUES (?, ?)", (KEY, json.dumps(items, ensure_ascii=False)))


def add_folder(con, path):
    if not path:
        raise LocalFileError("フォルダを選んでください")
    p = Path(path).expanduser()
    if not p.is_absolute() or not p.is_dir():
        raise LocalFileError("そのフォルダは見つかりません")
    p = p.resolve()
    if p.parent == p:
        raise LocalFileError("ドライブ全体は広すぎるので登録できません。中のフォルダを選んでください")
    items = folders(con)
    if any(Path(f).resolve() == p for f in items):
        raise LocalFileError("そのフォルダはもう入っています")
    _save(con, items + [str(p)])
    return folders(con)


def remove_folder(con, path):
    _save(con, [f for f in folders(con) if f != path])
    return folders(con)


def pick_folder():
    """Windows の「フォルダーの選択」の窓を出す(エクスプローラーと同じ窓)。選ばなければ None。
    Flask の中で窓を出すと固まることがあるので、別のプロセスで出す。"""
    code = ("import tkinter as tk\nfrom tkinter import filedialog\nr = tk.Tk(); r.withdraw(); r.attributes('-topmost', True)\n"
            "p = filedialog.askdirectory(title='モッチが探せるフォルダを選ぶ', mustexist=True)\nprint(p or '')")
    exe = sys.executable.replace("pythonw.exe", "python.exe")
    try:
        r = subprocess.run([exe, "-c", code], capture_output=True, text=True, encoding="utf-8", timeout=600,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except subprocess.TimeoutExpired:
        return None
    p = (r.stdout or "").strip()
    return str(Path(p)) if p else None


# ---------- 決まり ----------
def _hidden(p):
    if p.name.startswith("."):
        return True
    try:
        return bool(os.stat(p).st_file_attributes & stat.FILE_ATTRIBUTE_HIDDEN)
    except (AttributeError, OSError):
        return False


def _skip_dir(p):
    return p.name in SKIP_DIRS or _hidden(p) or p == DATA_DIR


def _secret_file(p):
    n = p.name.lower()
    return n == ".env" or n.startswith(".env.") or SECRET_NAME.search(n) is not None or p.suffix.lower() in KEY_EXT


def _place(con, path):
    """書かれた場所を、本当の場所にする。C:\\ から始まる場所か、登録したフォルダの名前から始まる場所(Workspace\\… など)。"""
    if not path:
        raise LocalFileError("ファイルの場所が要ります")
    p = Path(str(path).strip().strip('"').strip("'")).expanduser()
    if not p.is_absolute():
        parts = p.parts
        base = next((Path(f) for f in folders(con) if parts and Path(f).name.lower() == parts[0].lower()), None)
        if not base:
            raise LocalFileError("ファイルの場所は、file_search の答えにある場所をそのまま使ってください")
        p = base.joinpath(*parts[1:])
    return p.resolve()


def _inside(con, path, folder=False):
    """登録したフォルダの中のファイル(folder=True ならフォルダも)か確かめて、本当の場所を返す(.. やリンクで外に出られない)。"""
    p = _place(con, path)
    roots = [Path(f).resolve() for f in folders(con)]
    root = next((r for r in roots if p == r or r in p.parents), None)
    if not root:
        raise LocalFileError("そのファイルは、探してよいフォルダに入っていません。見てよければ、本人が設定の「フォルダ」に足します")
    for d in [p.parent, *p.parent.parents]:
        if d == root.parent:
            break
        if d != root and _skip_dir(d):
            raise LocalFileError("そのフォルダの中(隠しフォルダ・.git・段取りさんの data など)は読みません")
    if p == DATA_DIR or DATA_DIR in p.parents:
        raise LocalFileError("段取りさんの data フォルダは読みません")
    if folder and p.is_dir():
        if p != root and _skip_dir(p):
            raise LocalFileError("そのフォルダ(隠しフォルダ・.git など)は開きません")
        return p
    if _secret_file(p):
        raise LocalFileError("パスワードや鍵が入っていそうなファイルなので、読みません(規程第3条)")
    if not p.is_file():
        raise LocalFileError("そのファイルは見つかりません")
    return p


def _info(p, root):
    st = p.stat()
    return {"path": str(p), "name": p.name, "folder": str(p.parent.relative_to(root.parent)),
            "size_kb": round(st.st_size / 1024, 1), "modified": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M")}


# ---------- 探す・読む・開く ----------
def search(con, words, limit=SEARCH_LIMIT, ext=None):
    """名前(フォルダの名前も含む)に言葉が全部入るファイルを、新しい順に返す。中身は読まない。"""
    ws = [w.lower() for w in re.split(r"[\s ]+", words or "") if w]
    if not ws:
        raise LocalFileError("探す言葉が要ります")
    exts = {("." + e.lower().lstrip(".")) for e in (ext or [])}
    found, start, cut = [], time.monotonic(), False
    for f in folders(con):
        root = Path(f).resolve()
        if not root.is_dir():
            continue
        for d, dirs, files in os.walk(root):
            if time.monotonic() - start > SEARCH_SECONDS:
                cut = True
                break
            dp = Path(d)
            dirs[:] = [x for x in dirs if not _skip_dir(dp / x)]
            rel = str(dp.relative_to(root.parent)).lower()
            for n in files:
                p = dp / n
                if _secret_file(p) or _hidden(p) or (exts and p.suffix.lower() not in exts):
                    continue
                hay = rel + "\\" + n.lower()
                if all(w in hay for w in ws):
                    try:
                        found.append((p.stat().st_mtime, p, root))
                    except OSError:
                        continue
    found.sort(key=lambda x: x[0], reverse=True)
    return {"files": [_info(p, r) for _, p, r in found[:limit]], "more": max(0, len(found) - limit),
            "searched": folders(con), "stopped_early": cut}


def _decode(b):
    for enc in ("utf-8-sig", "cp932"):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue
    return b.decode("utf-8", errors="replace")


def _text_of(p):
    ext = p.suffix.lower()
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        out = [x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                out.append(" | ".join(c.text for c in row.cells))
        return "\n".join(out)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"## シート「{ws.title}」")
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    out.append(" | ".join(cells))
                if sum(len(x) for x in out) > MAX_CHARS:
                    break
        wb.close()
        return "\n".join(out)
    if ext == ".pptx":
        import pptx
        out = []
        for i, s in enumerate(pptx.Presentation(str(p)).slides, 1):
            out.append(f"## スライド {i}")
            out += [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        return "\n".join(out)
    if ext == ".pdf":
        import pypdf
        r = pypdf.PdfReader(str(p))
        out = []
        for i, pg in enumerate(r.pages, 1):
            out.append(f"## {i}ページ\n" + (pg.extract_text() or ""))
            if sum(len(x) for x in out) > MAX_CHARS:
                break
        return "\n".join(out)
    if ext in TEXT_EXT:
        return _decode(p.read_bytes()[: MAX_CHARS * 4])
    raise LocalFileError(f"{ext or '拡張子のない'}ファイルは、中身を文字にして読めません(名前・場所だけ分かります)")


def read(con, path):
    p = _inside(con, path)
    if p.stat().st_size > MAX_BYTES:
        raise LocalFileError("大きすぎるファイル(30MB より大きい)なので、読みません")
    try:
        text = _text_of(p)
    except LocalFileError:
        raise
    except Exception:
        raise LocalFileError("このファイルは開けませんでした(壊れている・パスワード付き・使っている途中 など)")
    root = next(Path(f).resolve() for f in folders(con) if Path(f).resolve() in p.parents)
    info = _info(p, root)
    info["text"] = text[:MAX_CHARS]
    info["cut"] = len(text) > MAX_CHARS
    return info


def check_open(con, path):
    p = _inside(con, path)
    if p.suffix.lower() in RUN_EXT:
        raise LocalFileError("動くファイル(.exe・.bat・.py など)は開きません")
    return p


def open_file(con, path, opener=None):
    """いつものアプリで開く(モッチからは確認のあと)。"""
    p = check_open(con, path)
    (opener or os.startfile)(str(p))
    return p


def open_link(con, path, opener=None):
    """モッチの答えの中の札を、本人が押したとき(S7-2 1.11)。ファイルはいつものアプリで、フォルダはエクスプローラーで開く。"""
    p = _inside(con, path, folder=True)
    if p.is_file() and p.suffix.lower() in RUN_EXT:
        raise LocalFileError("動くファイル(.exe・.bat・.py など)は開きません")
    (opener or os.startfile)(str(p))
    return p


def reveal(con, path, runner=None):
    """札の [フォルダ]。エクスプローラーを開き、そのファイルを選んだ形にする。"""
    p = _inside(con, path, folder=True)
    args = ["explorer", str(p)] if p.is_dir() else ["explorer", f"/select,{p}"]
    (runner or subprocess.Popen)(args)
    return p
